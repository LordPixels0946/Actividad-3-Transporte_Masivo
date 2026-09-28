# -*- coding: utf-8 -*-
"""
ANALIZADOR Y EXPORTADOR EXCEL
Genera reportes profesionales en formato Excel con análisis detallado
"""

import os
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, PieChart, Reference
import networkx as nx


class AnalizadorExcel:
    """Genera análisis profesional en Excel"""
    
    def __init__(self, base_conocimiento, output_dir='outputs'):
        self.base = base_conocimiento
        self.output_dir = output_dir
        self.grafo = self._crear_grafo()
        
        os.makedirs(output_dir, exist_ok=True)
    
    def _crear_grafo(self):
        """Crea grafo NetworkX"""
        G = nx.Graph()
        for estacion in self.base.estaciones.keys():
            G.add_node(estacion)
        for origen, destino, distancia, tiempo in self.base.conexiones:
            G.add_edge(origen, destino, weight=distancia, tiempo=tiempo)
        return G
    
    def generar_reporte_completo(self, resultados_busqueda=None, filename=None):
        """Genera reporte Excel completo con múltiples hojas"""
        if filename is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f'reporte_setp_{timestamp}.xlsx'
        
        filepath = os.path.join(self.output_dir, filename)
        
        # Crear Excel con escritor de pandas
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            # Hoja 1: Resumen de la red
            self._hoja_resumen_red(writer)
            
            # Hoja 2: Listado de estaciones
            self._hoja_estaciones(writer)
            
            # Hoja 3: Matriz de conexiones
            self._hoja_conexiones(writer)
            
            # Hoja 4: Análisis de centralidad
            self._hoja_centralidad(writer)
            
            # Hoja 5: Matriz de distancias
            self._hoja_matriz_distancias(writer)
            
            # Hoja 6: Resultados de búsquedas (si hay)
            if resultados_busqueda:
                self._hoja_resultados_busqueda(writer, resultados_busqueda)
            
            # Hoja 7: Estadísticas
            self._hoja_estadisticas(writer)
        
        # Aplicar formato profesional
        self._aplicar_formato_profesional(filepath)
        
        return filepath
    
    def _hoja_resumen_red(self, writer):
        """Hoja de resumen general"""
        data = {
            'Métrica': [
                'Total de Estaciones',
                'Total de Conexiones',
                'Conexiones Bidireccionales',
                'Distancia Total de Red (km)',
                'Distancia Promedio por Conexión (km)',
                'Grado Promedio',
                'Diámetro de la Red',
                'Radio de la Red',
                'Densidad de la Red',
                'Fecha de Análisis'
            ],
            'Valor': [
                len(self.grafo.nodes()),
                len(self.base.conexiones),
                'Sí',
                sum(d['weight'] for u, v, d in self.grafo.edges(data=True)),
                sum(d['weight'] for u, v, d in self.grafo.edges(data=True)) / len(self.base.conexiones),
                sum(dict(self.grafo.degree()).values()) / len(self.grafo.nodes()),
                nx.diameter(self.grafo) if nx.is_connected(self.grafo) else 'N/A',
                nx.radius(self.grafo) if nx.is_connected(self.grafo) else 'N/A',
                nx.density(self.grafo),
                datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            ]
        }
        
        df = pd.DataFrame(data)
        df.to_excel(writer, sheet_name='Resumen', index=False)
    
    def _hoja_estaciones(self, writer):
        """Hoja con detalle de cada estación"""
        data = []
        
        degree_centrality = nx.degree_centrality(self.grafo)
        betweenness = nx.betweenness_centrality(self.grafo, weight='weight')
        closeness = nx.closeness_centrality(self.grafo, distance='weight')
        
        for estacion in sorted(self.grafo.nodes()):
            coords = self.base.estaciones[estacion]
            vecinos = self.base.obtener_vecinos(estacion)
            
            data.append({
                'Estación': estacion,
                'Latitud': coords['lat'],
                'Longitud': coords['lon'],
                'Número de Conexiones': len(vecinos),
                'Centralidad de Grado': round(degree_centrality[estacion], 4),
                'Centralidad de Intermediación': round(betweenness[estacion], 4),
                'Centralidad de Cercanía': round(closeness[estacion], 4),
                'Vecinos': ', '.join([v[0] for v in vecinos])
            })
        
        df = pd.DataFrame(data)
        df.to_excel(writer, sheet_name='Estaciones', index=False)
    
    def _hoja_conexiones(self, writer):
        """Hoja con todas las conexiones"""
        data = []
        
        for origen, destino, distancia, tiempo in self.base.conexiones:
            data.append({
                'Origen': origen,
                'Destino': destino,
                'Distancia (km)': distancia,
                'Tiempo (min)': tiempo,
                'Velocidad Promedio (km/h)': round((distancia / tiempo) * 60, 2)
            })
        
        df = pd.DataFrame(data)
        df = df.sort_values('Distancia (km)', ascending=False)
        df.to_excel(writer, sheet_name='Conexiones', index=False)
    
    def _hoja_centralidad(self, writer):
        """Análisis de centralidad detallado"""
        degree_centrality = nx.degree_centrality(self.grafo)
        betweenness = nx.betweenness_centrality(self.grafo, weight='weight')
        closeness = nx.closeness_centrality(self.grafo, distance='weight')
        
        data = []
        for node in sorted(self.grafo.nodes()):
            importancia = (
                degree_centrality[node] * 0.4 +
                betweenness[node] * 0.4 +
                closeness[node] * 0.2
            )
            
            data.append({
                'Estación': node,
                'Centralidad de Grado': round(degree_centrality[node], 4),
                'Centralidad de Intermediación': round(betweenness[node], 4),
                'Centralidad de Cercanía': round(closeness[node], 4),
                'Puntuación de Importancia': round(importancia, 4)
            })
        
        df = pd.DataFrame(data)
        df = df.sort_values('Puntuación de Importancia', ascending=False)
        df['Ranking'] = range(1, len(df) + 1)
        df = df[['Ranking', 'Estación', 'Centralidad de Grado', 
                'Centralidad de Intermediación', 'Centralidad de Cercanía', 
                'Puntuación de Importancia']]
        
        df.to_excel(writer, sheet_name='Análisis Centralidad', index=False)
    
    def _hoja_matriz_distancias(self, writer):
        """Matriz de distancias entre todas las estaciones"""
        estaciones = sorted(self.grafo.nodes())
        matriz = {}
        
        for origen in estaciones:
            fila = {'Origen': origen}
            for destino in estaciones:
                if origen == destino:
                    fila[destino] = 0
                else:
                    try:
                        distancia = nx.shortest_path_length(
                            self.grafo, origen, destino, weight='weight')
                        fila[destino] = round(distancia, 2)
                    except nx.NetworkXNoPath:
                        fila[destino] = 'N/A'
            matriz[origen] = fila
        
        df = pd.DataFrame(matriz.values())
        df.to_excel(writer, sheet_name='Matriz Distancias', index=False)
    
    def _hoja_resultados_busqueda(self, writer, resultados):
        """Hoja con resultados de búsquedas realizadas"""
        data = []
        
        for i, resultado in enumerate(resultados, 1):
            data.append({
                'Búsqueda #': i,
                'Origen': resultado['origen'],
                'Destino': resultado['destino'],
                'Distancia (km)': resultado['distancia'],
                'Número de Paradas': resultado['paradas'],
                'Tiempo Estimado (min)': resultado['tiempo_estimado'],
                'Ruta Completa': ' → '.join(resultado['ruta']),
                'Fecha/Hora': resultado.get('timestamp', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
            })
        
        df = pd.DataFrame(data)
        df.to_excel(writer, sheet_name='Resultados Búsquedas', index=False)
    
    def _hoja_estadisticas(self, writer):
        """Estadísticas generales"""
        conexiones_distancias = [dist for o, d, dist, t in self.base.conexiones]
        conexiones_tiempos = [t for o, d, dist, t in self.base.conexiones]
        
        import numpy as np
        
        data = {
            'Estadística': [
                'Distancia Mínima entre Conexiones (km)',
                'Distancia Máxima entre Conexiones (km)',
                'Distancia Media (km)',
                'Desviación Estándar Distancia (km)',
                'Tiempo Mínimo entre Conexiones (min)',
                'Tiempo Máximo entre Conexiones (min)',
                'Tiempo Medio (min)',
                'Desviación Estándar Tiempo (min)',
                'Estación con Más Conexiones',
                'Estación con Menos Conexiones',
                'Número de Componentes Conexas',
                'Es Conexa la Red'
            ],
            'Valor': [
                min(conexiones_distancias),
                max(conexiones_distancias),
                round(np.mean(conexiones_distancias), 2),
                round(np.std(conexiones_distancias), 2),
                min(conexiones_tiempos),
                max(conexiones_tiempos),
                round(np.mean(conexiones_tiempos), 2),
                round(np.std(conexiones_tiempos), 2),
                max(self.grafo.degree(), key=lambda x: x[1])[0],
                min(self.grafo.degree(), key=lambda x: x[1])[0],
                nx.number_connected_components(self.grafo),
                'Sí' if nx.is_connected(self.grafo) else 'No'
            ]
        }
        
        df = pd.DataFrame(data)
        df.to_excel(writer, sheet_name='Estadísticas', index=False)
    
    def _aplicar_formato_profesional(self, filepath):
        """Aplica formato profesional al Excel"""
        wb = load_workbook(filepath)
        
        # Colores y estilos
        header_fill = PatternFill(start_color='2E86AB', end_color='2E86AB', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=12)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # Formato de encabezados
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border
            
            # Ajustar anchos de columna
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            # Aplicar bordes a todas las celdas con datos
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, 
                                   min_col=1, max_col=ws.max_column):
                for cell in row:
                    cell.border = border
                    if cell.row > 1:
                        cell.alignment = Alignment(horizontal='left', vertical='center')
        
        wb.save(filepath)
